"""Generate the customer-facing slide deck for the AI Lead Follow-Up Assistant.

Usage:
    python -m scripts.build_deck
Writes deliverables/AI_Lead_Follow-Up_Assistant_Slides.pptx
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUTPUT = Path(__file__).resolve().parent.parent / "deliverables" / "AI_Lead_Follow-Up_Assistant_Slides.pptx"

NAVY = RGBColor(0x0F, 0x17, 0x2A)
ACCENT = RGBColor(0x25, 0x63, 0xEB)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)
GRAY = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x16, 0x65, 0x34)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    return frame


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    _fill(bg, NAVY)

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.15), Inches(1.6), Inches(0.12))
    _fill(bar, ACCENT)

    frame = _textbox(slide, Inches(0.9), Inches(2.4), Inches(11.5), Inches(2.2))
    p = frame.paragraphs[0]
    run = p.add_run()
    run.text = "AI Lead Follow-Up Assistant"
    run.font.size = Pt(44)
    run.font.bold = True
    run.font.color.rgb = WHITE

    sub = frame.add_paragraph()
    sub.text = "Never lose another lead in your inbox."
    sub.font.size = Pt(22)
    sub.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)
    sub.space_before = Pt(12)

    tag = _textbox(slide, Inches(0.9), Inches(6.5), Inches(11.5), Inches(0.6))
    tp = tag.paragraphs[0]
    trun = tp.add_run()
    trun.text = "Human-in-the-loop · draft-only prototype"
    trun.font.size = Pt(14)
    trun.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    slide.notes_slide.notes_text_frame.text = (
        "We are not replacing your sales process. We make sure no lead slips through."
    )


def content_slide(prs, title, bullets, note=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(1.15))
    _fill(header, NAVY)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.15), SLIDE_W, Inches(0.08))
    _fill(accent, ACCENT)

    tf = _textbox(slide, Inches(0.6), Inches(0.22), Inches(12), Inches(0.75))
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = WHITE

    body = _textbox(slide, Inches(0.75), Inches(1.7), Inches(11.8), Inches(5.0))
    for index, (text, level) in enumerate(bullets):
        para = body.paragraphs[0] if index == 0 else body.add_paragraph()
        para.text = ("•  " if level == 0 else "–  ") + text
        para.level = level
        para.font.size = Pt(20 if level == 0 else 16)
        para.font.color.rgb = NAVY if level == 0 else GRAY
        para.space_after = Pt(12)

    if note:
        slide.notes_slide.notes_text_frame.text = note
    return slide


def flow_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(1.15))
    _fill(header, NAVY)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.15), SLIDE_W, Inches(0.08))
    _fill(accent, ACCENT)
    tf = _textbox(slide, Inches(0.6), Inches(0.22), Inches(12), Inches(0.75))
    run = tf.paragraphs[0].add_run()
    run.text = "How it works"
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = WHITE

    steps = [
        ("Import leads", "CSV or manual"),
        ("AI analyzes", "intent, priority, risk"),
        ("Today queue", "overdue, hot, stale"),
        ("Review rationale", "+ next action"),
        ("Draft reply", "you edit & copy"),
        ("Record outcome", "next date recalculated"),
    ]
    left = Inches(0.5)
    top = Inches(2.4)
    width = Inches(2.0)
    height = Inches(1.7)
    gap = Inches(0.12)
    for i, (head, sub) in enumerate(steps):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left + i * (width + gap), top, width, height
        )
        _fill(shape, LIGHT)
        shape.line.color.rgb = ACCENT
        frame = shape.text_frame
        frame.word_wrap = True
        p = frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = head
        r.font.bold = True
        r.font.size = Pt(14)
        r.font.color.rgb = NAVY
        p2 = frame.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = sub
        r2.font.size = Pt(11)
        r2.font.color.rgb = GRAY
        if i < len(steps) - 1:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                left + i * (width + gap) + width + Emu(20000),
                top + Inches(0.68),
                Inches(0.28),
                Inches(0.34),
            )
            _fill(arrow, ACCENT)

    caption = _textbox(slide, Inches(0.75), Inches(4.6), Inches(11.8), Inches(1.5))
    cp = caption.paragraphs[0]
    cr = cp.add_run()
    cr.text = "You stay in control at the draft step — every recommendation is explainable."
    cr.font.size = Pt(18)
    cr.font.italic = True
    cr.font.color.rgb = GRAY

    slide.notes_slide.notes_text_frame.text = "Each step is transparent; nothing is a black box."


def closing_slide(prs, title, bullets, note=""):
    slide = content_slide(prs, title, bullets, note)
    return slide


def build() -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    title_slide(prs)

    content_slide(
        prs,
        "The problem",
        [
            ("Leads arrive from web, referrals, email, events and ads — handled inconsistently.", 0),
            ("Owners ask the same question every morning: which leads need me today?", 0),
            ("Slow or forgotten follow-up is quiet, avoidable revenue loss.", 0),
        ],
        "Leads are easy to collect and easy to lose.",
    )

    content_slide(
        prs,
        "What it does",
        [
            ("Your daily answer: which leads need attention, why, and what to do next.", 0),
            ("Reads each lead and summarizes it in one line.", 0),
            ("Ranks High / Medium / Low with a plain-English reason.", 0),
            ("Recommends the next action and a follow-up date.", 0),
            ("Drafts a personalized reply you can send in seconds.", 0),
        ],
        "A decision-support copilot, not an auto-sender.",
    )

    flow_slide(prs)

    content_slide(
        prs,
        "What you see",
        [
            ("Today — the short list that matters: overdue, due today, new hot leads, needs review.", 0),
            ("Dashboard — your pipeline at a glance, filterable, counts that always match.", 0),
            ("Lead detail — raw lead, the AI's reasoning, the draft, and full history.", 0),
            ("Every recommendation shows a short why and a confidence level.", 0),
        ],
        "A user should know who to call within 10 seconds.",
    )

    content_slide(
        prs,
        "Built-in safety",
        [
            ("Draft-only. It never sends messages or makes commitments for you.", 0),
            ("No invented facts — no fake prices, availability, guarantees, or policies.", 0),
            ("Risk escalation — refunds, legal, complaints, sensitive data, regulated advice go to a human.", 0),
            ("Human approves everything before it leaves the building.", 0),
        ],
        "Guardrails are deterministic and tested, not left to the model.",
    )

    content_slide(
        prs,
        "Where we are / next",
        [
            ("Working prototype: import, analysis, priority queue, drafts, outcome tracking, dashboard.", 0),
            ("Runs in the cloud with sample data for a live demo.", 0),
            ("Next: connect your CRM/email/calendar so drafts flow into your tools.", 0),
            ("Next: durable storage and team roles for daily use.", 0),
        ],
        "We validated the workflow first; integrations follow real user demand.",
    )

    closing_slide(
        prs,
        "Success metrics (backup)",
        [
            ("95%+ of leads receive valid, structured analysis.", 0),
            ("Users agree with the priority on 80%+ of scenarios.", 0),
            ("70%+ of drafts need only light edits.", 0),
            ("Zero fabricated prices, policies, or commitments in testing.", 0),
        ],
        "These targets come from the prototype test plan.",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT))
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
