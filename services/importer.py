"""CSV/XLSX lead import with validation and per-row error reporting.

Implements S1-01: valid rows import with unique lead IDs; invalid rows are
reported without losing valid rows; required columns are identified clearly.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import IO, Any, Dict, List, Optional, Union

from models.schemas import Lead, ResponseStatus, Stage, new_id

CANONICAL_COLUMNS = {
    "leadid": "lead_id",
    "id": "lead_id",
    "name": "name",
    "company": "company",
    "email": "email",
    "phone": "phone",
    "leadsource": "source",
    "source": "source",
    "inquiry": "inquiry",
    "message": "inquiry",
    "question": "inquiry",
    "details": "inquiry",
    "datereceived": "received_at",
    "received": "received_at",
    "receivedat": "received_at",
    "lastcontactdate": "last_contact_at",
    "lastcontact": "last_contact_at",
    "lastcontactat": "last_contact_at",
    "leadstage": "stage",
    "stage": "stage",
    "priorityseed": "priority_seed",
    "priority": "priority_seed",
    "intentseed": "intent_seed",
    "intent": "intent_seed",
    "nextactionseed": "next_action_seed",
    "nextaction": "next_action_seed",
    "followupdate": "follow_up_date",
    "responsestatus": "response_status",
    "status": "response_status",
    "notes": "notes",
    "scenariotag": "scenario_tag",
    "scenario": "scenario_tag",
}

REQUIRED_COLUMNS = ("name", "inquiry", "stage", "received_at")
CRITICAL_LABELS = {
    "name": "Name",
    "inquiry": "Inquiry",
    "stage": "Lead Stage",
    "received_at": "Date Received",
}

EXCEL_EPOCH = datetime(1899, 12, 30)


def _normalise_key(key: str) -> str:
    return "".join(ch for ch in str(key).lower() if ch.isalnum())


def _map_headers(headers: List[str]) -> Dict[str, str]:
    mapping: Dict[str, str] = {}
    for header in headers:
        canonical = CANONICAL_COLUMNS.get(_normalise_key(header))
        if canonical and canonical not in mapping:
            mapping[canonical] = header
    return mapping


def parse_date(value: Any) -> Optional[datetime]:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, (int, float)):
        try:
            return EXCEL_EPOCH + timedelta(days=float(value))
        except (OverflowError, ValueError):
            return None
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%Y-%m-%dT%H:%M:%S", "%m/%d/%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        return None


def _parse_stage(value: Any) -> Stage:
    return Stage(str(value).strip())


def _parse_response_status(value: Any) -> ResponseStatus:
    text = str(value).strip() if value not in (None, "") else "Not Contacted"
    return ResponseStatus(text)


@dataclass
class ImportRowError:
    row: int
    message: str


@dataclass
class ImportReport:
    imported: int = 0
    failed: int = 0
    errors: List[ImportRowError] = field(default_factory=list)
    missing_columns: List[str] = field(default_factory=list)
    lead_ids: List[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return self.failed == 0 and not self.missing_columns

    def summary(self) -> str:
        parts = [f"{self.imported} imported", f"{self.failed} rejected"]
        if self.missing_columns:
            parts.append("missing columns: " + ", ".join(self.missing_columns))
        return "; ".join(parts)


def _read_rows(source: Union[str, IO], filename: Optional[str] = None) -> List[Dict[str, Any]]:
    name = (filename or (source if isinstance(source, str) else "")).lower()
    if name.endswith(".csv"):
        if isinstance(source, str):
            with open(source, newline="", encoding="utf-8-sig") as handle:
                return list(csv.DictReader(handle))
        raw = source.read()
        if isinstance(raw, bytes):
            raw = raw.decode("utf-8-sig")
        return list(csv.DictReader(io.StringIO(raw)))
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - dependency guard
            raise RuntimeError("openpyxl is required to import .xlsx files") from exc
        if isinstance(source, str):
            workbook = load_workbook(source, read_only=True, data_only=True)
        else:
            source.seek(0)
            workbook = load_workbook(io.BytesIO(source.read()), read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        rows = list(sheet.iter_rows(values_only=True))
        workbook.close()
        if not rows:
            return []
        headers = [str(h) if h is not None else "" for h in rows[0]]
        records = []
        for values in rows[1:]:
            if values is None or all(v is None for v in values):
                continue
            records.append(dict(zip(headers, values)))
        return records
    raise ValueError(f"Unsupported file type: {filename or source}")


def import_leads(
    source: Union[str, IO],
    repository,
    filename: Optional[str] = None,
) -> ImportReport:
    """Import leads from a CSV/XLSX source into the repository."""
    rows = _read_rows(source, filename)
    report = ImportReport()

    if not rows:
        report.errors.append(ImportRowError(0, "No data rows found"))
        return report

    header_mapping = _map_headers(list(rows[0].keys()))
    missing = [CRITICAL_LABELS[c] for c in REQUIRED_COLUMNS if c not in header_mapping]
    if missing:
        report.missing_columns = missing
        report.errors.append(
            ImportRowError(1, "Missing required columns: " + ", ".join(missing))
        )
        return report

    seen_ids: set[str] = set()
    for index, raw_row in enumerate(rows, start=2):  # row 1 is the header
        mapped = {
            canonical: raw_row.get(header)
            for canonical, header in header_mapping.items()
        }
        try:
            lead = _build_lead(mapped, seen_ids, row_number=index)
        except ValueError as exc:
            report.failed += 1
            report.errors.append(ImportRowError(index, str(exc)))
            continue
        seen_ids.add(lead.lead_id)
        repository.upsert_lead(lead)
        report.imported += 1
        report.lead_ids.append(lead.lead_id)

    return report


def _build_lead(mapped: Dict[str, Any], seen_ids: set, row_number: int) -> Lead:
    name = (mapped.get("name") or "").strip() if isinstance(mapped.get("name"), str) else mapped.get("name")
    inquiry = (mapped.get("inquiry") or "").strip() if isinstance(mapped.get("inquiry"), str) else mapped.get("inquiry")

    if not name:
        raise ValueError("Missing required field: Name")
    if not inquiry:
        raise ValueError("Missing required field: Inquiry")

    stage_raw = mapped.get("stage")
    if stage_raw in (None, ""):
        raise ValueError("Missing required field: Lead Stage")
    try:
        stage = _parse_stage(stage_raw)
    except ValueError:
        raise ValueError(f"Invalid Lead Stage: {stage_raw!r}")

    received = parse_date(mapped.get("received_at"))
    if received is None:
        raise ValueError("Missing or invalid Date Received")

    last_contact = parse_date(mapped.get("last_contact_at"))

    try:
        response_status = _parse_response_status(mapped.get("response_status"))
    except ValueError:
        raise ValueError(f"Invalid Response Status: {mapped.get('response_status')!r}")

    lead_id = mapped.get("lead_id")
    if lead_id not in (None, ""):
        lead_id = str(lead_id).strip()
        if lead_id in seen_ids:
            raise ValueError(f"Duplicate Lead ID: {lead_id}")
    else:
        lead_id = new_id("lead")

    return Lead(
        lead_id=lead_id,
        name=str(name),
        company=_clean(mapped.get("company")),
        email=_clean(mapped.get("email")),
        phone=_clean(mapped.get("phone")),
        source=_clean(mapped.get("source")) or "Unknown",
        inquiry=str(inquiry),
        stage=stage,
        received_at=received,
        last_contact_at=last_contact,
        response_status=response_status,
        notes=_clean(mapped.get("notes")),
    )


def _clean(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None
